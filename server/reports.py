"""Сбор данных за смену и генерация отчётов PDF / Excel / Word."""
from __future__ import annotations
import asyncio
import io
import time
from xml.sax.saxutils import escape
from . import case_data, db, shifts, ai
from .config import FONT_DIR
from .sim import LAY

DEV = {d["id"]: d for d in LAY["devices"]}
ST = {"OK": "Норма", "WARNING": "Предупр.", "CRITICAL": "Авария", "INFO": "Инфо", "STOPPED": "Остановлен"}
PR = {"high": "Высокий", "med": "Средний", "low": "Низкий"}
OS = {"new": "Новый", "work": "В работе", "done": "Выполнен"}
FORMATS = {"pdf": "PDF", "xlsx": "Excel", "docx": "Word"}
CAP = 600  # макс. «вес» одной записи телеметрии (сек), чтобы пропуски в данных не считались работой


# ---------------------------------------------------------------- данные
def shift_metrics(start: float, end: float):
    out = []
    for d in LAY["devices"]:
        rows = db.rows("SELECT ts,temp,vib,units,status FROM snapshots WHERE device_id=? AND ts>=? AND ts<? ORDER BY ts",
                       (d["id"], int(start), int(end)))
        t = {"OK": 0.0, "WARNING": 0.0, "CRITICAL": 0.0, "STOPPED": 0.0}
        units, tmax, vmax = 0.0, 0.0, 0.0
        for i, r in enumerate(rows):
            if i + 1 < len(rows):
                w = min(rows[i + 1]["ts"] - r["ts"], CAP)
                du = rows[i + 1]["units"] - r["units"]
                if 0 <= du < 1e6:
                    units += du
            else:
                w = min(end - r["ts"], 60)
            t[r["status"]] = t.get(r["status"], 0.0) + max(0, w)
            tmax, vmax = max(tmax, r["temp"]), max(vmax, r["vib"])
        cov = sum(t.values())
        plan = d["plan_per_hour"] * cov / 3600
        pct = units / plan if plan else 0.0
        stop = t["CRITICAL"] + t["STOPPED"]      # аварии и остановки считаются простоем
        a = 1 - stop / cov if cov else 0
        p = min(1.0, pct)
        q = (t["OK"] * 0.99 + t["WARNING"] * 0.95 + t["CRITICAL"] * 0.8 + t["STOPPED"] * 0.99) / cov if cov else 0
        out.append({"id": d["id"], "name": d["name"], "shop": d["shop"], "units": round(units), "plan": round(plan),
                    "pct": pct, "a": a, "p": p, "q": q, "oee": a * p * q, "cov": cov,
                    "down_min": round(stop / 60, 1), "warn_min": round(t["WARNING"] / 60, 1),
                    "tmax": tmax, "vmax": vmax, "status": rows[-1]["status"] if rows else "—"})
    return out


def build(which: str = "cur") -> dict:
    w = shifts.window(which)
    devs = shift_metrics(w["start"], w["until"])
    n = max(1, len(devs))
    units, plan = sum(x["units"] for x in devs), sum(x["plan"] for x in devs)
    inc = db.rows("SELECT ts,level,device_id,message FROM incidents WHERE ts>=? AND ts<? ORDER BY ts", (int(w["start"]), int(w["until"])))
    orders = db.rows("SELECT * FROM orders ORDER BY id")
    now = time.time()
    o_rows, done_in = [], 0
    for o in orders:
        in_win = o["done_ts"] and w["start"] <= o["done_ts"] < w["until"]
        if in_win:
            done_in += 1
        if o["status"] != "done" or in_win or (o["created_ts"] and w["start"] <= o["created_ts"] < w["until"]):
            o_rows.append({"id": o["id"], "title": o["title"], "device": DEV.get(o["device_id"], {}).get("name", o["device_id"]),
                           "assignee": o["assignee"], "priority": PR.get(o["priority"], o["priority"]),
                           "status": OS.get(o["status"], o["status"]), "due": o["due"].replace("T", " ") if o["due"] else "",
                           "late": o["status"] != "done" and bool(o["due"]) and o["due"] < time.strftime("%Y-%m-%dT%H:%M", time.localtime(now))})
    rep = {
        "title": "Отчёт за смену", "shift": w["name"], "which": which,
        "start": shifts.fmt(w["start"]), "end": shifts.fmt(w["until"]), "generated": shifts.fmt(now, "%d.%m.%Y %H:%M:%S"),
        "partial": which == "cur" and not w.get("gap"), "date_key": shifts.fmt(w["start"], "%Y-%m-%d"), "idx": w["idx"],
        "source": "симулятор (демо-данные)",
        "totals": {"units": units, "plan": plan, "pct": units / plan if plan else 0,
                   "oee": sum(x["oee"] for x in devs) / n, "a": sum(x["a"] for x in devs) / n,
                   "p": sum(x["p"] for x in devs) / n, "q": sum(x["q"] for x in devs) / n,
                   "down_min": round(sum(x["down_min"] for x in devs), 1), "warn_min": round(sum(x["warn_min"] for x in devs), 1),
                   "alerts": sum(1 for i in inc if i["level"] in ("WARNING", "CRITICAL")),
                   "crit": sum(1 for i in inc if i["level"] == "CRITICAL"),
                   "orders_done": done_in, "orders_open": sum(1 for o in orders if o["status"] != "done"),
                   "cov_h": round(max([x["cov"] for x in devs] + [0]) / 3600, 1), "win_h": round((w["until"] - w["start"]) / 3600, 1)},
        "devices": devs,
        "incidents": [{"time": shifts.fmt(i["ts"], "%d.%m %H:%M"), "level": ST.get(i["level"], i["level"]),
                       "device": DEV.get(i["device_id"], {}).get("name", "—"), "message": i["message"]} for i in inc[-40:]],
        "incidents_total": len(inc), "orders": o_rows, "ai": None, "case": None,
    }
    t = rep["totals"]
    rep["case"] = case_data.summary({"oee": t["oee"], "defect": 1 - t["q"], "downtime_today": {x["id"]: x["down_min"] for x in devs}})
    return rep


def brief(rep: dict) -> str:
    """Компактный текст для ИИ."""
    t = rep["totals"]
    lines = [f"{rep['shift']}, {rep['start']} — {rep['end']}" + (" (смена ещё идёт)" if rep["partial"] else ""),
             f"Выпуск {t['units']} из плана {t['plan']} ({t['pct']*100:.0f}%), OEE {t['oee']*100:.0f}%, простои {t['down_min']} мин, предупреждения {t['warn_min']} мин, "
             f"тревог {t['alerts']} (аварий {t['crit']}), нарядов закрыто {t['orders_done']}, открыто {t['orders_open']}. Данные за {t['cov_h']} из {t['win_h']} ч."]
    for d in rep["devices"]:
        lines.append(f"- {d['name']} ({d['shop']}): выпуск {d['units']}/{d['plan']} ({d['pct']*100:.0f}%), OEE {d['oee']*100:.0f}%, простой {d['down_min']} мин, "
                     f"пик T {d['tmax']:.0f}°C, вибрация {d['vmax']:.3f}g, сейчас {ST.get(d['status'], d['status'])}")
    for i in rep["incidents"][-8:]:
        lines.append(f"* {i['time']} {i['level']}: {i['message']}")
    c = rep.get("case")
    if c:
        lines.append("Цели кейса: " + "; ".join(f"{k}: цель {g}, факт {f} ({st})" for k, g, f, st in _case_checks(rep)))
        lines += ["Отклонения по тестовым данным кейса: " + "; ".join(c["alerts"])] if c["alerts"] else []
    for o in rep["orders"]:
        lines.append(f"# наряд #{o['id']} {o['title']} — {o['status']}{' (просрочен)' if o['late'] else ''}")
    return "\n".join(lines)


async def add_ai(rep: dict):
    if not ai.configured():
        return rep
    try:
        rep["ai"] = await ai.complete(
            "Ты аналитик производства автозавода. По данным смены напиши резюме для руководителя: 3–5 коротких предложений "
            "(итог, главные проблемы, что стоит сделать). Только по данным, без выдумок, без markdown, по-русски.",
            [{"role": "user", "content": brief(rep)}], max_tokens=500, timeout=40)
    except Exception:
        rep["ai"] = None
    return rep


def context_text(sim) -> str:
    """Срез состояния завода для ИИ-чата."""
    parts = [f"Сейчас: {shifts.fmt(time.time(), '%d.%m.%Y %H:%M')}. Источник данных: симулятор (демо)."]
    m = sim.last
    if m:
        s = m["stats"]
        parts.append(f"Живые показатели: OEE {s['oee']*100:.0f}%, доступность {s['availability']*100:.0f}%, производительность {s['performance']*100:.0f}%, качество {s['quality']*100:.0f}%.")
        for d in m["devices"]:
            x = DEV[d["device_id"]]
            line = f"- {x['name']} ({x['shop']}): {ST[d['status']]}, T={d['temperature']}°C, вибрация {d['vibration']}g, загрузка {d['load_percentage']}%"
            if d.get("ai"):
                line += f", ПРОГНОЗ: {d['ai']['message']}"
            parts.append(line)
        parts.append("Последние события: " + "; ".join(f"{i['time']} {i['type']} {i['message']}" for i in m["incidents"][:6]))
    parts.append("Итоги текущей смены:\n" + brief(build("cur")))
    parts.append(case_data.text(sim))
    return "\n".join(parts)


# ---------------------------------------------------------------- рендер
def _pct(v):
    return f"{v*100:.0f}%"


def _case_checks(rep):
    """Сверка показателей смены с целями кейса: [(показатель, цель, факт, статус)]."""
    c, T = rep["case"], rep["case"]["targets"]
    v, f = c["live"], c["forecast"]
    ok = lambda b: "Норма" if b else "Отклонение"
    return [("OEE", f"не менее {T['oee']*100:.0f}%", _pct(v["oee"]), ok(v["oee_ok"])),
            ("Брак (оценка по симулятору)", f"не более {T['defect_max']*100:.0f}%", f"{v['defect']*100:.1f}%", ok(v["defect_ok"])),
            ("Макс. простой станка за смену, мин", f"не более {T['downtime_max_min']} в сутки", v["downtime_max"], ok(v["downtime_ok"])),
            ("Выпуск за месяц, годных (прогноз по данным завода)", f"не менее {T['month_plan_min']}", f['month_good'], ok(f["month_good"] >= T["month_plan_min"]))]


def _case_tables(rep):
    """Таблицы приложения «Показатели кейса»: [(заголовок, шапка, строки)]."""
    c = rep["case"]
    return [
        ("Выполнение плана по линиям (данные завода из загруженного файла)", ["Линия", "План", "Факт", "% плана", "Загрузка, %", "Годные / план"],
         [[x["line"], x["plan"], x["fact"], _pct(x["pct"]), f"{x['load']:.0f}", _pct(x["oee"])] for x in c["by_line"]]),
        ("Качество по участкам (допустимый брак ≤ 2%)", ["Дата", "Участок", "Выпущено", "Брак", "% брака", "Статус"],
         [[x["date"], x["section"], x["out"], x["defect"], f"{x['pct']*100:.1f}%", "Норма" if x["ok"] else "Выше нормы"] for x in c["quality"]]),
        ("Простои оборудования (лимит 60 мин в сутки)", ["Дата", "Участок", "Оборудование", "Причина", "Мин", "Тип", "% лимита"],
         [[x["date"], x["section"], x["equipment"], x["reason"], x["min"], "плановый" if x["planned"] else "внеплановый", _pct(x["share"])] for x in c["downtime"]]),
        ("План выпуска на месяц", ["Модель", "План, авто", "Доля"],
         [[m["model"], m["plan"], _pct(m["share"])] for m in c["models"]] + [["Итого", c["plan_sum"], "100%"]]),
    ]


def _rows_devices(rep):
    return [[d["name"], d["shop"], d["units"], d["plan"], _pct(d["pct"]), _pct(d["oee"]), d["down_min"], d["warn_min"],
             f"{d['tmax']:.0f}", f"{d['vmax']:.3f}", ST.get(d["status"], d["status"])] for d in rep["devices"]]


DEV_HEAD = ["Оборудование", "Цех", "Выпуск", "План", "% плана", "OEE", "Простой, мин", "Предупр., мин", "T макс, °C", "Вибр. макс, g", "Сейчас"]


def _kpis(rep):
    t = rep["totals"]
    return [("Выпуск (факт / план)", f"{t['units']} / {t['plan']}  ({_pct(t['pct'])})"), ("OEE", _pct(t["oee"])),
            ("Доступность / Производительность / Качество", f"{_pct(t['a'])} / {_pct(t['p'])} / {_pct(t['q'])}"),
            ("Простои (аварии), мин", t["down_min"]), ("Время в предупреждении, мин", t["warn_min"]),
            ("Тревог / из них аварий", f"{t['alerts']} / {t['crit']}"), ("Нарядов закрыто за смену / открыто сейчас", f"{t['orders_done']} / {t['orders_open']}"),
            ("Данные получены за", f"{t['cov_h']} ч из {t['win_h']} ч")]


def _sub(rep):
    return f"{rep['shift']} · {rep['start']} — {rep['end']}" + (" (смена ещё идёт)" if rep["partial"] else "")


def filename(rep, fmt):
    return f"Allur_report_{rep['date_key']}_shift{rep['idx']}{'_partial' if rep['partial'] else ''}.{fmt}"


def render(rep: dict, fmt: str):
    fn = {"pdf": _pdf, "xlsx": _xlsx, "docx": _docx}[fmt]
    return filename(rep, fmt), fn(rep)


async def make_files(which, fmts, with_ai=True):
    rep = await asyncio.to_thread(build, which)
    if with_ai:
        await add_ai(rep)
    files = [await asyncio.to_thread(render, rep, f) for f in fmts]
    return rep, files


def _pdf(rep):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    if "DV" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DV", str(FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DVB", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
    G = colors.HexColor("#0B8F5E")
    h1 = ParagraphStyle("h1", fontName="DVB", fontSize=17, leading=22, textColor=G, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName="DVB", fontSize=11.5, spaceBefore=10, spaceAfter=4)
    tx = ParagraphStyle("tx", fontName="DV", fontSize=9, leading=12)
    sm = ParagraphStyle("sm", fontName="DV", fontSize=7.5, leading=9.5)
    smb = ParagraphStyle("smb", fontName="DVB", fontSize=7.5, leading=9.5, textColor=colors.white)

    def table(head, rows, widths, zebra=True):
        data = [[Paragraph(escape(str(h)), smb) for h in head]] + [[Paragraph(escape(str(c)), sm) for c in r] for r in rows]
        t = Table(data, colWidths=widths, repeatRows=1)
        st = [("BACKGROUND", (0, 0), (-1, 0), G), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C9D3CE")),
              ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        if zebra:
            st += [("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6F4")])]
        t.setStyle(TableStyle(st))
        return t

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"{rep['title']} — {rep['shift']}", author="Allur Digital Twin")
    W = doc.width
    el = [Paragraph("Allur Digital Twin · Отчёт за смену", h1), Paragraph(escape(_sub(rep)), tx),
          Paragraph(f"Сформирован: {rep['generated']} · Источник данных: {escape(rep['source'])}", sm), Spacer(1, 6)]
    el.append(table(["Показатель", "Значение"], [[k, v] for k, v in _kpis(rep)], [W * 0.5, W * 0.5]))
    el += [Paragraph("Сверка с целями кейса", h2)]
    el.append(table(["Показатель", "Цель", "Факт за смену", "Статус"], [list(r) for r in _case_checks(rep)], [W * x for x in (0.42, 0.22, 0.2, 0.16)]))
    if rep.get("ai"):
        el += [Paragraph("Резюме ИИ", h2), Paragraph(escape(rep["ai"]).replace("\n", "<br/>"), tx)]
    el += [Paragraph("Оборудование", h2)]
    wd = [0.17, 0.12, 0.07, 0.07, 0.07, 0.07, 0.09, 0.09, 0.08, 0.09, 0.08]
    el.append(table(DEV_HEAD, _rows_devices(rep), [W * x for x in wd]))
    el += [Paragraph(f"События за смену (всего {rep['incidents_total']}, показаны последние {len(rep['incidents'])})", h2)]
    el.append(table(["Время", "Уровень", "Оборудование", "Описание"], [[i["time"], i["level"], i["device"], i["message"]] for i in rep["incidents"]]
                    or [["—", "", "", "Событий не было"]], [W * 0.1, W * 0.1, W * 0.18, W * 0.62]))
    el += [Paragraph("Наряды", h2)]
    el.append(table(["№", "Задача", "Оборудование", "Исполнитель", "Приоритет", "Статус", "Срок"],
                    [[o["id"], o["title"], o["device"], o["assignee"], o["priority"], o["status"] + (" · просрочен" if o["late"] else ""), o["due"]] for o in rep["orders"]]
                    or [["—", "Нарядов нет", "", "", "", "", ""]], [W * x for x in (0.05, 0.27, 0.16, 0.18, 0.09, 0.14, 0.11)]))
    h2s = ParagraphStyle("h2s", parent=h2, spaceBefore=5, spaceAfter=2)
    el += [PageBreak(), Paragraph("Приложение · Показатели кейса", h1)]
    for ttl, head, rows in _case_tables(rep):
        el += [Paragraph(escape(ttl), h2s), table(head, rows, [W / len(head)] * len(head))]
    if rep["case"]["alerts"]:
        el += [Paragraph("Отклонения и риски", h2s)] + [Paragraph("• " + escape(a), sm) for a in rep["case"]["alerts"]]
    doc.build(el)
    return buf.getvalue()


def _xlsx(rep):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

    wb = Workbook()
    F, FB = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_fill, head_font = PatternFill("solid", fgColor="0B8F5E"), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    thin = Side(style="thin", color="C9D3CE")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet(ws, head, rows, widths):
        for c, h in enumerate(head, 1):
            x = ws.cell(row=1, column=c, value=h)
            x.font, x.fill, x.border = head_font, head_fill, box
            x.alignment = Alignment(wrap_text=True, vertical="center")
        for r, row in enumerate(rows, 2):
            for c, v in enumerate(row, 1):
                x = ws.cell(row=r, column=c, value=v)
                x.font, x.border = F, box
                x.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Сводка"
    ws["A1"], ws["A1"].font = "Allur Digital Twin · Отчёт за смену", Font(name="Arial", size=14, bold=True, color="0B8F5E")
    ws["A2"], ws["A2"].font = _sub(rep), F
    ws["A3"], ws["A3"].font = f"Сформирован: {rep['generated']} · Источник: {rep['source']}", Font(name="Arial", size=9, color="697872")
    for i, (k, v) in enumerate(_kpis(rep), 5):
        ws.cell(row=i, column=1, value=k).font = FB
        ws.cell(row=i, column=2, value=v).font = F
    r = 5 + len(_kpis(rep)) + 1
    if rep.get("ai"):
        ws.cell(row=r, column=1, value="Резюме ИИ").font = FB
        c = ws.cell(row=r + 1, column=1, value=rep["ai"])
        c.font, c.alignment = F, Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=2)
        ws.row_dimensions[r + 1].height = 110
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 44

    wd = wb.create_sheet("Оборудование")
    rows = [[d["name"], d["shop"], d["units"], d["plan"], round(d["pct"], 4), round(d["oee"], 4), d["down_min"], d["warn_min"],
             round(d["tmax"], 1), d["vmax"], ST.get(d["status"], d["status"])] for d in rep["devices"]]
    sheet(wd, DEV_HEAD, rows, [24, 20, 10, 10, 10, 10, 13, 14, 12, 13, 12])
    for r in range(2, len(rows) + 2):
        wd.cell(row=r, column=5).number_format = wd.cell(row=r, column=6).number_format = "0%"
    ch = BarChart()
    ch.type, ch.title, ch.height, ch.width = "col", "Выпуск: факт и план", 8, 18
    ch.add_data(Reference(wd, min_col=3, max_col=4, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
    ch.set_categories(Reference(wd, min_col=1, min_row=2, max_row=len(rows) + 1))
    wd.add_chart(ch, f"A{len(rows) + 4}")

    sheet(wb.create_sheet("События"), ["Время", "Уровень", "Оборудование", "Описание"],
          [[i["time"], i["level"], i["device"], i["message"]] for i in rep["incidents"]], [14, 12, 24, 80])
    sheet(wb.create_sheet("Наряды"), ["№", "Задача", "Оборудование", "Исполнитель", "Приоритет", "Статус", "Срок", "Просрочен"],
          [[o["id"], o["title"], o["device"], o["assignee"], o["priority"], o["status"], o["due"], "да" if o["late"] else ""] for o in rep["orders"]],
          [6, 36, 24, 24, 12, 12, 18, 11])
    sheet(wb.create_sheet("Цели кейса"), ["Показатель", "Цель", "Факт за смену", "Статус"], [list(r) for r in _case_checks(rep)], [52, 24, 18, 14])
    wc = wb.create_sheet("Показатели кейса")
    r0 = 1
    for ttl, head, rows in _case_tables(rep):
        wc.cell(row=r0, column=1, value=ttl).font = FB
        for c, h in enumerate(head, 1):
            x = wc.cell(row=r0 + 1, column=c, value=h)
            x.font, x.fill, x.border = head_font, head_fill, box
        for i, row in enumerate(rows, r0 + 2):
            for c, v in enumerate(row, 1):
                x = wc.cell(row=i, column=c, value=v)
                x.font, x.border = F, box
        r0 += len(rows) + 4
    wc.cell(row=r0, column=1, value="Отклонения и риски").font = FB
    for i, a in enumerate(rep["case"]["alerts"], r0 + 1):
        wc.cell(row=i, column=1, value="• " + a).font = F
    for i, w in enumerate([22, 14, 16, 22, 12, 14, 12]):
        wc.column_dimensions[chr(65 + i)].width = w
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _docx(rep):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt, Mm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = Document()
    s = d.sections[0]
    s.orientation, s.page_width, s.page_height = WD_ORIENT.LANDSCAPE, Mm(297), Mm(210)
    s.left_margin = s.right_margin = Mm(14)
    s.top_margin = s.bottom_margin = Mm(12)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    GREEN = RGBColor(0x0B, 0x8F, 0x5E)

    def heading(text, size):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold, r.font.size, r.font.color.rgb, r.font.name = True, Pt(size), GREEN, "Arial"
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(3)

    def shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"), sh.set(qn("w:color"), "auto"), sh.set(qn("w:fill"), color)
        tcPr.append(sh)

    def table(head, rows, size=8, w=None):
        t = d.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        t.autofit = False
        for i, h in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(str(h))
            r.bold, r.font.size, r.font.color.rgb = True, Pt(size), RGBColor(255, 255, 255)
            shade(c, "0B8F5E")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(size)
        if w:
            for i, col in enumerate(t.columns):
                col.width = Mm(w[i])
            for row in t.rows:
                for i, c in enumerate(row.cells):
                    c.width = Mm(w[i])
        return t

    heading("Allur Digital Twin · Отчёт за смену", 18)
    d.add_paragraph(_sub(rep))
    p = d.add_paragraph(f"Сформирован: {rep['generated']} · Источник данных: {rep['source']}")
    p.runs[0].font.size = Pt(8)
    table(["Показатель", "Значение"], [[k, v] for k, v in _kpis(rep)], 9, w=[120, 90])
    heading("Сверка с целями кейса", 12)
    table(["Показатель", "Цель", "Факт за смену", "Статус"], [list(r) for r in _case_checks(rep)], 9, w=[110, 60, 45, 36])
    if rep.get("ai"):
        heading("Резюме ИИ", 12)
        for line in rep["ai"].split("\n"):
            if line.strip():
                d.add_paragraph(line.strip())
    heading("Оборудование", 12)
    table(DEV_HEAD, _rows_devices(rep), w=[40, 36, 17, 17, 17, 17, 21, 24, 20, 24, 36])
    heading(f"События за смену (всего {rep['incidents_total']}, показаны последние {len(rep['incidents'])})", 12)
    table(["Время", "Уровень", "Оборудование", "Описание"], [[i["time"], i["level"], i["device"], i["message"]] for i in rep["incidents"]] or [["—", "", "", "Событий не было"]], w=[28, 24, 48, 169])
    heading("Наряды", 12)
    table(["№", "Задача", "Оборудование", "Исполнитель", "Приоритет", "Статус", "Срок"],
          [[o["id"], o["title"], o["device"], o["assignee"], o["priority"], o["status"] + (" · просрочен" if o["late"] else ""), o["due"]] for o in rep["orders"]]
          or [["—", "Нарядов нет", "", "", "", "", ""]], w=[12, 70, 40, 46, 24, 40, 37])
    d.add_page_break()
    heading("Приложение · Показатели кейса", 16)
    for ttl, head, rows in _case_tables(rep):
        heading(ttl, 12)
        table(head, rows, 9)
    if rep["case"]["alerts"]:
        heading("Отклонения и риски", 12)
        for a in rep["case"]["alerts"]:
            d.add_paragraph("• " + a)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()
