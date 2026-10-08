"""Загрузка данных завода из файла docx: таблицы «работа линий», «простои», «план по моделям», «показатели качества»,
схема участков и дополнительные вводные (из текста). Файл из папки data/import/ подхватывается автоматически при первом
запуске; любой другой файл такой же структуры можно загрузить кнопкой на сайте (POST /api/import).
Данные хранятся в базе (таблица kv) и заменяют встроенные значения в case_data.py."""
from __future__ import annotations
import json
import re
import time
from . import db
from .config import DATA_DIR

IMPORT_DIR = DATA_DIR / "import"
KEY = "plant_data"


def _num(v, default=0.0):
    try:
        return float(str(v).strip().replace("\xa0", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return default


def _kind(head: list[str]) -> str | None:
    h = " ".join(head).lower()
    if "линия" in h and "факт" in h:
        return "lines"
    if "оборудование" in h and "причина" in h:
        return "downtime"
    if "модель" in h:
        return "models"
    if "брак" in h and "участок" in h:
        return "quality"
    return None


def parse_docx(src) -> dict:
    """src — путь или файловый объект. Возвращает словарь с таблицами (пустые, если не найдены)."""
    from docx import Document
    doc = Document(src)
    out = {"lines": [], "downtime": [], "models": [], "quality": [], "scheme": [], "targets": {}}
    for t in doc.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        if len(rows) < 2:
            continue
        k = _kind(rows[0])
        for r in rows[1:]:
            if not any(r):
                continue
            if k == "lines" and len(r) >= 6:
                out["lines"].append([r[0], r[1], int(_num(r[2])), int(_num(r[3])), _num(r[4]), int(_num(r[5]))])
            elif k == "downtime" and len(r) >= 5:
                out["downtime"].append([r[0], r[1], r[2], r[3], int(_num(r[4])), bool(re.search(r"планов", r[3], re.I))])
            elif k == "models" and len(r) >= 2:
                out["models"].append([r[0], int(_num(r[1]))])
            elif k == "quality" and len(r) >= 4:
                out["quality"].append([r[0], r[1], int(_num(r[2])), int(_num(r[3]))])
    text = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for line in text:
        if "→" in line:
            out["scheme"] = [x.strip() for x in line.split("→") if x.strip()]
        m = re.search(r"(\d+)\s+смен\w*\s+по\s+(\d+)\s+час", line)
        if m:
            out["targets"]["shifts"], out["targets"]["shift_hours"] = int(m.group(1)), int(m.group(2))
        m = re.search(r"OEE[^\d]*(\d+(?:[.,]\d+)?)\s*%", line, re.I)
        if m:
            out["targets"]["oee"] = _num(m.group(1)) / 100
        m = re.search(r"брак\w*[^\d]*(\d+(?:[.,]\d+)?)\s*%", line, re.I)
        if m:
            out["targets"]["defect_max"] = _num(m.group(1)) / 100
        m = re.search(r"простой[^\d]*(\d+)\s*мин", line, re.I)
        if m:
            out["targets"]["downtime_max_min"] = int(m.group(1))
        m = re.search(r"не менее\s+([\d\s\xa0]+)\s*автомоб", line, re.I)
        if m:
            out["targets"]["month_plan_min"] = int(_num(m.group(1)))
    return out


def counts(d: dict) -> dict:
    return {k: len(d.get(k) or []) for k in ("lines", "downtime", "models", "quality")}


def save(data: dict, source: str) -> dict:
    data = {**data, "source": source, "ts": int(time.time())}
    db.kv_set(KEY, json.dumps(data, ensure_ascii=False))
    try:
        from . import orders
        orders.sync_docx()          # задачи из файла -> «Панель мастера»
    except Exception as e:
        print(f"  [!] Наряды из файла не созданы: {e}")
    return status()


def load() -> dict | None:
    try:
        raw = db.kv_get(KEY)
        return json.loads(raw) if raw else None
    except ValueError:
        return None


def status() -> dict:
    d = load()
    if not d:
        return {"loaded": False}
    return {"loaded": True, "source": d.get("source", ""), "ts": d.get("ts", 0), "counts": counts(d), "targets": d.get("targets", {})}


def ensure_loaded():
    """При первом запуске подхватывает первый .docx из data/import/."""
    if load():
        return
    try:
        for f in sorted(IMPORT_DIR.glob("*.docx")):
            data = parse_docx(str(f))
            if any(counts(data).values()):
                save(data, f.name)
                print(f"  Данные завода загружены из {f.name}: {counts(data)}")
                return
    except Exception as e:
        print(f"  [!] Не удалось загрузить данные из data/import: {e}")
